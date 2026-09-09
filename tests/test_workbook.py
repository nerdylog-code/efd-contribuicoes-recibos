from pathlib import Path

from openpyxl import load_workbook

from efd_receipts.parser import parse_receipt_text
from efd_receipts.workbook import HEADERS, write_workbook


def test_workbook_contains_headers_and_status(tmp_path: Path) -> None:
    row = parse_receipt_text("CNPJ: 00.000.000/0000-00", "partial.txt")
    output = write_workbook([row], tmp_path / "receipts.xlsx")
    workbook = load_workbook(output, data_only=True)
    sheet = workbook["Recibos"]
    assert [cell.value for cell in sheet[2]] == HEADERS
    assert sheet.cell(3, 8).value == "INCOMPLETE"
    workbook.close()
