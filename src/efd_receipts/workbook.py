from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .parser import ReceiptData

HEADERS = [
    "Arquivo",
    "CNPJ",
    "Numero_Recibo",
    "Data_SERPRO",
    "Hora_SERPRO",
    "Agente_Receptor",
    "Periodo_Apuracao",
    "Status",
]


def write_workbook(rows: list[ReceiptData], output_path: Path) -> Path:
    """Write normalized receipt rows to a new workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Recibos"
    sheet.append(["EFD Contribuicoes - recibos"])
    sheet.append(HEADERS)
    for item in rows:
        sheet.append([
            item.source_file,
            item.cnpj,
            item.receipt_number,
            item.received_date,
            item.received_time,
            item.receiving_agent,
            item.period,
            item.status,
        ])

    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
    for cell in sheet[2]:
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(HEADERS))}{max(2, sheet.max_row)}"
    for column in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column) + 2
        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(width, 40)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
